"""Tests for extract_error_report_1xxx.py."""
from __future__ import annotations

import csv
import sys

import pytest
from openpyxl import load_workbook

from conftest import load_script

mod = load_script("extract_error_report_1xxx.py")

COLS = ["householdId", "deviceId", "session.clientIp", "sessionId", "content_uuid", "start_time", "error_codes",
        "error_descriptions", "state_sequence", "min_position", "max_position", "position_not_moved", "raw_user_agent"]


def write_source(root, tenant, iso, codes):
    d = root / tenant
    d.mkdir(parents=True, exist_ok=True)
    with (d / f"{iso}_setupsession_with_outcome.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        for i, c in enumerate(codes):
            w.writerow({"sessionId": f"s{i}", "error_codes": c, "position_not_moved": "false", "raw_user_agent": "ua"})


def run(monkeypatch, root, *args):
    monkeypatch.setattr(sys, "argv", ["x", "--output-dir", str(root), *args])
    mod.main()


def test_row_has_1xxx_excludes_6xxx_and_other_codes():
    assert mod.row_has_1xxx({"error_codes": "1003"})
    assert mod.row_has_1xxx({"error_codes": "6001;1002"})
    assert not mod.row_has_1xxx({"error_codes": "6001"})
    assert not mod.row_has_1xxx({"error_codes": "3016, E31012, 10030"})
    assert not mod.row_has_1xxx({"error_codes": ""})


def test_filter_rows_keeps_columns_and_not_only_not_moved():
    rows = [{"sessionId": "a", "error_codes": "1003", "position_not_moved": "false", "raw_user_agent": "ua"},
            {"sessionId": "b", "error_codes": "6001"}]
    out = mod.filter_rows(rows, ["sessionId", "error_codes", "position_not_moved"])
    assert out == [{"sessionId": "a", "error_codes": "1003", "position_not_moved": "false"}]


def test_parse_tenants_dedupes_uppercases_and_rejects_empty():
    assert mod.parse_tenants("za, GH,za") == ["ZA", "GH"]
    with pytest.raises(ValueError):
        mod.parse_tenants(" , ")


def test_single_date_writes_one_csv_per_tenant(tmp_path, monkeypatch):
    for t in ("ZA", "GH"):
        write_source(tmp_path, t, "2026-10-08", ["1003", "6001", ""])
    run(monkeypatch, tmp_path, "--date", "08102026", "--tenants", "ZA,GH")
    for t in ("ZA", "GH"):
        rows = list(csv.DictReader((tmp_path / t / "2026-10-08_error_report_1xxx.csv").open()))
        assert [r["sessionId"] for r in rows] == ["s0"]


def test_range_writes_one_xlsx_per_tenant_with_sheet_per_date(tmp_path, monkeypatch):
    for t in ("ZA", "ZM"):
        for iso in ("2026-10-07", "2026-10-08"):
            write_source(tmp_path, t, iso, ["1003"])
    run(monkeypatch, tmp_path, "--from", "07102026", "--to", "08102026", "--tenants", "ZA,ZM")
    files = sorted((tmp_path / "ZA").glob("*error_report_1xxx.xlsx"))
    assert len(files) == 1 and len(list((tmp_path / "ZM").glob("*error_report_1xxx.xlsx"))) == 1
    assert load_workbook(files[0]).sheetnames == ["2026-10-07", "2026-10-08"]


def test_missing_tenant_input_exits_nonzero_naming_path_and_writes_nothing(tmp_path, monkeypatch, capsys):
    write_source(tmp_path, "ZA", "2026-10-08", ["1003"])
    with pytest.raises(SystemExit) as exc:
        run(monkeypatch, tmp_path, "--date", "08102026", "--tenants", "ZA,GH")
    assert exc.value.code == 1
    assert "GH" in capsys.readouterr().err
    assert not list((tmp_path / "ZA").glob("*error_report_1xxx*"))


def test_reversed_range_exits_nonzero(tmp_path, monkeypatch):
    with pytest.raises(SystemExit) as exc:
        run(monkeypatch, tmp_path, "--from", "09102026", "--to", "08102026")
    assert exc.value.code == 1
