"""Tests for combine_error_code_by_opco.py."""
from __future__ import annotations

import csv
import importlib.util
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

_SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))
_SPEC = importlib.util.spec_from_file_location("ctap_combine_opco", _SCRIPTS_DIR / "combine_error_code_by_opco.py")
assert _SPEC and _SPEC.loader
mod = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(mod)


def _flat(root: Path, tenant: str, name: str, header: list[str], rows: list[list[str]]) -> None:
    d = root / tenant / mod.FLAT_DIR
    d.mkdir(parents=True, exist_ok=True)
    with open(d / name, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def _run(monkeypatch, *argv):
    monkeypatch.setattr(sys, "argv", ["x", *argv])
    monkeypatch.setattr(mod, "init_logging", lambda *_: None)
    mod.main()


def test_combines_tenants_dates_and_codes_skipping_all(tmp_path, monkeypatch) -> None:
    _flat(tmp_path, "GH", "2026-10-08_6007.csv", ["sessionId", "errorcode"], [["a", "6007"]])
    _flat(tmp_path, "GH", "2026-10-09_1003.csv", ["sessionId", "errorcode", "extra"], [["b", "1003", "x"]])
    _flat(tmp_path, "GH", "2026-10-09_all.csv", ["sessionId", "errorcode"], [["dup", "1003"]])
    _flat(tmp_path, "ZM", "2026-10-11_6007.csv", ["sessionId", "errorcode"], [["c", "6007"]])
    (tmp_path / "SA").mkdir()  # no flat folder -> not a tenant
    out = tmp_path / "combined.xlsx"
    _run(monkeypatch, "--output-dir", str(tmp_path), "--out", str(out))
    wb = load_workbook(out)
    assert wb.sheetnames == ["GH", "ZM"]
    gh = [[c.value for c in r] for r in wb["GH"].iter_rows()]
    assert gh[0] == ["date", "sessionId", "errorcode", "extra"]
    assert [r[1] for r in gh[1:]] == ["a", "b"]
    assert gh[2][0] == "2026-10-09"


def test_date_range_filters_files(tmp_path, monkeypatch) -> None:
    _flat(tmp_path, "GH", "2026-10-07_6007.csv", ["sessionId"], [["old"]])
    _flat(tmp_path, "GH", "2026-10-08_6007.csv", ["sessionId"], [["in"]])
    _run(monkeypatch, "--output-dir", str(tmp_path), "--from", "08102026", "--to", "11102026")
    out = tmp_path / "2026-10-08_2026-10-11_error_by_opco.xlsx"
    rows = [[c.value for c in r] for r in load_workbook(out)["GH"].iter_rows()]
    assert rows == [["date", "sessionId"], ["2026-10-08", "in"]]


def test_empty_range_gives_header_only_sheet(tmp_path, monkeypatch) -> None:
    _flat(tmp_path, "GH", "2026-10-07_6007.csv", ["sessionId"], [["old"]])
    out = tmp_path / "e.xlsx"
    _run(monkeypatch, "--output-dir", str(tmp_path), "--from", "01112026", "--out", str(out))
    assert [[c.value for c in r] for r in load_workbook(out)["GH"].iter_rows()] == [["date"]]


def test_reversed_range_and_no_tenants_fail(tmp_path, monkeypatch) -> None:
    with pytest.raises(SystemExit):
        _run(monkeypatch, "--output-dir", str(tmp_path), "--from", "11102026", "--to", "08102026")
    with pytest.raises(SystemExit):
        _run(monkeypatch, "--output-dir", str(tmp_path))


def test_list_code_files_ignores_non_matching(tmp_path) -> None:
    _flat(tmp_path, "GH", "2026-10-08_E31012.csv", ["a"], [["1"]])
    (tmp_path / "GH" / mod.FLAT_DIR / "notes.txt").write_text("x")
    files = mod.list_code_files(str(tmp_path / "GH" / mod.FLAT_DIR), date(2026, 10, 8), None)
    assert [(i, c) for i, c, _ in files] == [("2026-10-08", "E31012")]
