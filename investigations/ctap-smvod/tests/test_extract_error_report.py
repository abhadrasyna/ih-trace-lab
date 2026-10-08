"""Tests for extract_error_report_1xxx_6xxx.py."""
from __future__ import annotations

import csv
import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

_SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))
_SPEC = importlib.util.spec_from_file_location("ctap_extract_err", _SCRIPTS_DIR / "extract_error_report_1xxx_6xxx.py")
assert _SPEC and _SPEC.loader
mod = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(mod)

SOURCE_COLS = [c for c in mod.DEFAULT_COLUMNS if c != "is_wv"] + ["raw_user_agent", "extra"]
WV_UA = "Mozilla/5.0 (Linux; Android 15; SM-A055F Build/AP3A; wv) AppleWebKit/537.36"
PLAIN_UA = "Mozilla/5.0 (Linux; Android 15; SM-A055F Build/AP3A) AppleWebKit/537.36"


def _write_source(out_dir: Path, iso: str, codes: list[str]) -> None:
    with open(out_dir / f"{iso}_setupsession_with_outcome.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=SOURCE_COLS)
        w.writeheader()
        for i, c in enumerate(codes):
            w.writerow({"sessionId": f"s{i}", "error_codes": c, "extra": "x"})


def _run(monkeypatch, *argv):
    monkeypatch.setattr(sys, "argv", ["x", *argv])
    monkeypatch.setattr(mod, "init_logging", lambda *_: None)
    mod.main()


def test_row_filter_matches_exact_1xxx_6xxx_only() -> None:
    assert mod.row_has_1xxx_6xxx({"error_codes": "3016; 6001"})
    assert mod.row_has_1xxx_6xxx({"error_codes": "1003"})
    assert not mod.row_has_1xxx_6xxx({"error_codes": "E31012; 3016; 16001"})
    assert not mod.row_has_1xxx_6xxx({"error_codes": ""})


def test_single_date_writes_filtered_csv_with_position_report_columns(tmp_path, monkeypatch) -> None:
    _write_source(tmp_path, "2026-10-06", ["6001", "3016", "1003; 3014", ""])
    _run(monkeypatch, "--date", "06102026", "--output-dir", str(tmp_path))
    rows = list(csv.DictReader(open(tmp_path / "2026-10-06_error_report_1xxx_6xxx.csv")))
    assert [r["sessionId"] for r in rows] == ["s0", "s2"]
    assert list(rows[0].keys()) == mod.DEFAULT_COLUMNS


def test_range_writes_xlsx_sheet_per_date(tmp_path, monkeypatch) -> None:
    _write_source(tmp_path, "2026-10-05", ["6002"])
    _write_source(tmp_path, "2026-10-06", ["3016"])
    _run(monkeypatch, "--from", "05102026", "--to", "06102026", "--output-dir", str(tmp_path))
    out = next(tmp_path.glob("*error_report_1xxx_6xxx.xlsx"))
    wb = load_workbook(out)
    assert wb.sheetnames == ["2026-10-05", "2026-10-06"]
    assert wb["2026-10-05"].max_row == 2 and wb["2026-10-06"].max_row == 1


def test_range_missing_date_aborts_without_output(tmp_path, monkeypatch, capsys) -> None:
    _write_source(tmp_path, "2026-10-05", ["6002"])
    with pytest.raises(SystemExit):
        _run(monkeypatch, "--from", "05102026", "--to", "06102026", "--output-dir", str(tmp_path))
    assert "06102026" in capsys.readouterr().err
    assert not list(tmp_path.glob("*.xlsx"))


def test_wv_user_agent_marked_true() -> None:
    rows = mod.filter_rows([{"error_codes": "6001", "raw_user_agent": WV_UA}], mod.DEFAULT_COLUMNS)
    assert rows[0]["is_wv"] == "true"


def test_non_wv_user_agent_marked_false() -> None:
    rows = mod.filter_rows([{"error_codes": "6001", "raw_user_agent": PLAIN_UA}], mod.DEFAULT_COLUMNS)
    assert rows[0]["is_wv"] == "false"


def test_missing_user_agent_is_blank_not_false() -> None:
    rows = mod.filter_rows([{"error_codes": "6001", "raw_user_agent": ""}], mod.DEFAULT_COLUMNS)
    assert rows[0]["is_wv"] == ""


def test_is_wv_column_in_xlsx(tmp_path, monkeypatch) -> None:
    src = tmp_path / "2026-10-05_setupsession_with_outcome.csv"
    with open(src, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=SOURCE_COLS)
        w.writeheader()
        w.writerow({"sessionId": "s0", "error_codes": "6002", "raw_user_agent": WV_UA})
    _run(monkeypatch, "--from", "05102026", "--to", "05102026", "--output-dir", str(tmp_path))
    ws = load_workbook(next(tmp_path.glob("*error_report_1xxx_6xxx.xlsx")))["2026-10-05"]
    header = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=header.index("is_wv") + 1).value == "true"


def test_xlsx_keeps_equals_prefixed_text_as_string(tmp_path) -> None:
    out = tmp_path / "o.xlsx"
    mod.write_xlsx(str(out), ["a"], {"s": [{"a": "=1+1"}]})
    assert load_workbook(out)["s"]["A2"].data_type == "s"


def test_source_without_error_codes_column_raises(tmp_path) -> None:
    p = tmp_path / "s.csv"
    p.write_text("sessionId\ns1\n")
    with pytest.raises(ValueError):
        mod.read_source(str(p))


def test_reversed_range_rejected() -> None:
    with pytest.raises(ValueError):
        mod.dates_in_range("06102026", "05102026")
