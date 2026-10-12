#!/usr/bin/env python3
"""
Combine the per-error-code flat files (output/<TENANT>/error_eventdata_flat/
<date>_<errorcode>.csv, written by flatten_error_eventdata.py
--split-by-errorcode) from all opcos into ONE .xlsx PER ERROR CODE (flattened
columns differ by code, so codes are never mixed): one sheet per opco, rows from
all dates stacked, with a leading `date` column. An opco with no rows for a
code gets no sheet in that code's file. The `<date>_all.csv` files are ignored.

Tenants are every sub-folder of the output dir that has an error_eventdata_flat
folder. The date range and --codes are optional; omit them to use every date
and every code found. Reads existing files only -- never fires Athena/Lightstep.

Usage:
    python3 combine_error_code_by_opco.py                              # all dates, all codes
    python3 combine_error_code_by_opco.py --from 08102026 --to 11102026 --codes 6007,1003
    # -> output/error_code/6007/2026-10-08_2026-10-11_error_6007_by_opco.xlsx, output/error_code/1003/...
"""
import argparse
import logging
import os
import re
import sys
from datetime import date

from lib import csv_io_compat as csv
from lib.bootstrap import REPO_ROOT  # noqa: F401  # import side effect: sys.path bootstrap
from src.lib.logging.bootstrap import init_logging

from lib.runtime import ddmmyyyy_to_date, investigation_output_dir

LOGGER = logging.getLogger("investigations.ctap_smvod.scripts.combine_error_code_by_opco")

FLAT_DIR = "error_eventdata_flat"
CODE_DIR = "error_code"
FILE_RE = re.compile(r"^(\d{4}-\d{2}-\d{2})_(?!all\.csv$)(.+)\.csv$")


def discover_tenants(output_dir: str) -> list[str]:
    """Sorted sub-folders of output_dir that contain an error_eventdata_flat folder."""
    return sorted(
        d for d in os.listdir(output_dir)
        if os.path.isdir(os.path.join(output_dir, d, FLAT_DIR))
    )


def list_code_files(flat_dir: str, start: date | None, end: date | None) -> list[tuple[str, str, str]]:
    """(iso_date, errorcode, path) for per-code files inside [start, end], sorted by date then code."""
    found = []
    for name in os.listdir(flat_dir):
        m = FILE_RE.match(name)
        if not m:
            continue
        iso, code = m.groups()
        d = date.fromisoformat(iso)
        if (start and d < start) or (end and d > end):
            continue
        found.append((iso, code, os.path.join(flat_dir, name)))
    return sorted(found)


def read_rows(path: str) -> list[dict]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def combine_tenant(files: list[tuple[str, str, str]]) -> tuple[list[str], list[dict]]:
    """Stack rows of all files, prefixing a `date` column; columns are the ordered union."""
    columns: list[str] = ["date"]
    rows: list[dict] = []
    for iso, _code, path in files:
        for r in read_rows(path):
            for c in r:
                if c not in columns:
                    columns.append(c)
            rows.append({"date": iso, **r})
    return columns, rows


def write_xlsx(path: str, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, (columns, rows) in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.append(columns)
        for r in rows:
            ws.append([r.get(c, "") for c in columns])
            for cell in ws[ws.max_row]:
                # telemetry text starting with '=' must stay text, not a formula
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


def default_out_name(code: str, start: date | None, end: date | None) -> str:
    prefix = f"{start.isoformat()}_{end.isoformat()}_" if start and end else (
        f"from_{start.isoformat()}_" if start else f"until_{end.isoformat()}_" if end else "all_dates_")
    return f"{prefix}error_{code}_by_opco.xlsx"


def group_by_code(files: list[tuple[str, str, str]]) -> dict[str, list[tuple[str, str, str]]]:
    grouped: dict[str, list[tuple[str, str, str]]] = {}
    for f in files:
        grouped.setdefault(f[1], []).append(f)
    return grouped


def main():
    init_logging("ctap-smvod-combine_error_code_by_opco")
    LOGGER.info("Starting combine_error_code_by_opco")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--from", dest="date_from", help="Optional range start DDMMYYYY (inclusive)")
    ap.add_argument("--to", dest="date_to", help="Optional range end DDMMYYYY (inclusive)")
    ap.add_argument("--output-dir", default=None, help="Root holding <TENANT>/ folders (default: investigation output dir)")
    ap.add_argument("--codes", default=None, help="Optional comma-separated error codes (default: every code found)")
    args = ap.parse_args()

    out_dir = str(investigation_output_dir(args.output_dir))
    try:
        start = ddmmyyyy_to_date(args.date_from) if args.date_from else None
        end = ddmmyyyy_to_date(args.date_to) if args.date_to else None
        if start and end and start > end:
            raise ValueError(f"--from {args.date_from} is after --to {args.date_to}")
        tenants = discover_tenants(out_dir)
        if not tenants:
            raise ValueError(f"no <TENANT>/{FLAT_DIR} folders under {out_dir}")
    except (ValueError, OSError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)

    wanted = {c.strip() for c in args.codes.split(",") if c.strip()} if args.codes else None
    by_tenant = {
        t: group_by_code([f for f in list_code_files(os.path.join(out_dir, t, FLAT_DIR), start, end)
                          if wanted is None or f[1] in wanted])
        for t in tenants
    }
    codes = sorted({c for grouped in by_tenant.values() for c in grouped})
    if not codes:
        print("ERROR: no per-code files found for the requested tenants/dates/codes", file=sys.stderr)
        sys.exit(1)
    for missing in sorted((wanted or set()) - set(codes)):
        LOGGER.warning("No files found for requested code %s", missing)

    for code in codes:
        sheets = {t: combine_tenant(g[code]) for t, g in by_tenant.items() if code in g}
        out_path = os.path.join(out_dir, CODE_DIR, code, default_out_name(code, start, end))
        write_xlsx(out_path, sheets)
        summary = ", ".join(f"{t}={len(r)}" for t, (_c, r) in sheets.items())
        print(f"Wrote {out_path} (rows per opco: {summary})", file=sys.stderr)
        LOGGER.info("Wrote error code %s out=%s %s", code, out_path, summary)
    LOGGER.info("Completed combine_error_code_by_opco codes=%d", len(codes))


if __name__ == "__main__":
    main()
