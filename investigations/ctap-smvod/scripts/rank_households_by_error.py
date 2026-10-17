#!/usr/bin/env python3
"""
Rank households by number of error sessions in a `*_error_report_1xxx*.xlsx` (one sheet per date).
Per household: distinct sessions, distinct days, distinct devices, error codes seen. Reads the xlsx
Fails loudly on a missing column (householdId, sessionId, deviceId, error_codes) or blank householdId/sessionId.

Usage:
    python3 rank_households_by_error.py --xlsx ../output/ZA/2026-10-01_2026-10-13_error_report_1xxx.xlsx --top 25
    # -> stdout table + <xlsx stem>_household_ranking.csv next to the xlsx (or in --out-dir)
"""
import argparse
import csv
import logging
import os
import re
import sys

from lib.bootstrap import REPO_ROOT  # noqa: F401  # import side effect: sys.path bootstrap
from src.lib.logging.bootstrap import init_logging

LOGGER = logging.getLogger("investigations.ctap_smvod.scripts.rank_households_by_error")

REQUIRED = ("householdId", "sessionId", "deviceId", "error_codes")
COLUMNS = ["rank", "householdId", "sessions", "days", "devices", "error_codes"]


def load_rows(path: str) -> list[dict]:
    """Read every sheet; each row gets a `date` key from its sheet name."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    rows: list[dict] = []
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if header is None:
                continue
            missing = [c for c in REQUIRED if c not in header]
            if missing:
                raise ValueError(f"sheet {ws.title!r} missing column(s): {', '.join(missing)}")
            for vals in it:
                if any(v is not None for v in vals):
                    rows.append(dict(zip(header, vals)) | {"date": ws.title})
    finally:
        wb.close()
    return rows


def rank_households(rows: list[dict]) -> list[dict]:
    """Pure: one dict per household, sorted by distinct sessions desc then householdId."""
    acc: dict[str, dict] = {}
    for r in rows:
        hh = str(r.get("householdId") or "").strip()
        sid = str(r.get("sessionId") or "").strip()
        if not hh or not sid:
            raise ValueError(f"blank householdId/sessionId on a row (date {r.get('date')})")
        a = acc.setdefault(hh, {"sessions": set(), "days": set(), "devices": set(), "codes": set()})
        a["sessions"].add(sid)
        a["days"].add(r["date"])
        if r.get("deviceId"):
            a["devices"].add(str(r["deviceId"]))
        a["codes"].update(c.strip() for c in re.split(r"[,;]", str(r.get("error_codes") or "")) if c.strip())
    ranked = sorted(acc.items(), key=lambda kv: (-len(kv[1]["sessions"]), kv[0]))
    return [{"rank": i, "householdId": hh, "sessions": len(a["sessions"]), "days": len(a["days"]),
             "devices": len(a["devices"]), "error_codes": ",".join(sorted(a["codes"]))}
            for i, (hh, a) in enumerate(ranked, 1)]


def main() -> None:
    init_logging("ctap-smvod-rank_households_by_error")
    LOGGER.info("Starting rank_households_by_error")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, help="Path to *_error_report_1xxx*.xlsx")
    ap.add_argument("--top", type=int, default=25, help="Rows to print to stdout (CSV always has all)")
    ap.add_argument("--out-dir", default=None, help="Directory for the CSV (default: next to the xlsx)")
    args = ap.parse_args()
    if args.top < 0:
        ap.error("--top must be >= 0")
    try:
        rows = load_rows(args.xlsx)
        ranking = rank_households(rows)
    except (ValueError, FileNotFoundError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)

    out_dir = args.out_dir or os.path.dirname(os.path.abspath(args.xlsx))
    os.makedirs(out_dir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(args.xlsx))[0]
    out_path = os.path.join(out_dir, f"{stem}_household_ranking.csv")
    with open(out_path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(ranking)

    total = sum(r["sessions"] for r in ranking)
    print(f"{len(rows)} rows, {total} sessions, {len(ranking)} households")
    print("\t".join(COLUMNS))
    for r in ranking[: args.top]:
        print("\t".join(str(r[c]) for c in COLUMNS))
    print(f"wrote {out_path}", file=sys.stderr)
    LOGGER.info("Completed rank_households_by_error households=%d out=%s", len(ranking), out_path)


if __name__ == "__main__":
    main()
