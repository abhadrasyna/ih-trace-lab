#!/usr/bin/env python3
"""
Extract every session whose error_codes contain a 1xxx or 6xxx code (exact
4-digit tokens, regex ^[16]\\d{3}$; e.g. 1003, 6001 -- but not 3016 or
E31012) from output/YYYY-MM-DD_setupsession_with_outcome.csv, keeping the same
columns as extract_position_report.py's position_report.csv. Unlike that
report, rows are NOT limited to position_not_moved.

Reads existing pipeline outputs only -- never fires Athena/Lightstep. If an
input is missing, aborts naming the dates to run through run_pipeline.py.

Usage:
    # one date -> output/YYYY-MM-DD_error_report_1xxx_6xxx.csv
    python3 extract_error_report_1xxx_6xxx.py --date 06102026

    # date range -> one .xlsx, one sheet per date
    python3 extract_error_report_1xxx_6xxx.py --from 01102026 --to 06102026
"""
import argparse
import logging
import os
import re
import sys
from datetime import timedelta

from lib import csv_io_compat as csv
from lib.bootstrap import REPO_ROOT  # noqa: F401  # import side effect: sys.path bootstrap
from src.lib.logging.bootstrap import init_logging

from lib.runtime import (
    ddmmyyyy_to_date,
    investigation_output_dir,
    range_output_path,
    snapshot_output_path,
)
from lib.webview import is_wv_cell

LOGGER = logging.getLogger("investigations.ctap_smvod.scripts.extract_error_report_1xxx_6xxx")

DEFAULT_COLUMNS = [
    "householdId", "deviceId", "session.clientIp", "sessionId", "content_uuid",
    "start_time", "error_codes", "error_descriptions", "state_sequence",
    "min_position", "max_position", "position_not_moved", "is_wv",
]
ERROR_CODE_RE = re.compile(r"^[16]\d{3}$")
ARTIFACT = "error_report_1xxx_6xxx"


def row_has_1xxx_6xxx(row: dict) -> bool:
    """True if any ';'/','-separated token of the row's error_codes cell is
    an exact 4-digit code starting with 1 or 6."""
    cell = row.get("error_codes", "") or ""
    tokens = (t.strip() for chunk in cell.split(";") for t in chunk.split(","))
    return any(ERROR_CODE_RE.match(t) for t in tokens)


def filter_rows(source_rows: list[dict], columns: list[str]) -> list[dict]:
    return [
        {c: r.get(c, "") for c in columns} | ({"is_wv": is_wv_cell(r)} if "is_wv" in columns else {})
        for r in source_rows if row_has_1xxx_6xxx(r)
    ]


def dates_in_range(start_ddmmyyyy: str, end_ddmmyyyy: str) -> list[str]:
    """Inclusive list of DDMMYYYY dates.

    Raises:
        ValueError: If a date is malformed or start is after end.
    """
    start, end = ddmmyyyy_to_date(start_ddmmyyyy), ddmmyyyy_to_date(end_ddmmyyyy)
    if start > end:
        raise ValueError(f"--from {start_ddmmyyyy} is after --to {end_ddmmyyyy}")
    return [(start + timedelta(days=i)).strftime("%d%m%Y") for i in range((end - start).days + 1)]


def read_source(path: str) -> list[dict]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        missing = [c for c in ("sessionId", "error_codes") if c not in (reader.fieldnames or [])]
        if missing:
            raise ValueError(f"column(s) {missing} missing in {path} (error data not merged yet?)")
        return list(reader)


def write_csv(path: str, columns: list[str], rows: list[dict]) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: str, columns: list[str], sheets: dict[str, list[dict]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        ws.append(columns)
        for r in rows:
            ws.append([r.get(c, "") for c in columns])
            for cell in ws[ws.max_row]:
                # telemetry text starting with '=' must stay text, not a formula
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


def main():
    init_logging("ctap-smvod-extract_error_report_1xxx_6xxx")
    LOGGER.info("Starting extract_error_report_1xxx_6xxx")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--date", help="Single date DDMMYYYY -> CSV output")
    ap.add_argument("--from", dest="date_from", help="Range start DDMMYYYY (with --to) -> .xlsx, one sheet per date")
    ap.add_argument("--to", dest="date_to", help="Range end DDMMYYYY, inclusive")
    ap.add_argument("--output-dir", default=None, help="Directory holding inputs/outputs (default: investigation output dir)")
    ap.add_argument("--out", default=None, help="Output path override (.csv for --date, .xlsx for a range)")
    ap.add_argument("--columns", default=",".join(DEFAULT_COLUMNS), help="Comma-separated columns to keep")
    args = ap.parse_args()

    if bool(args.date) == bool(args.date_from or args.date_to) or bool(args.date_from) != bool(args.date_to):
        ap.error("pass either --date, or both --from and --to")

    out_dir = str(investigation_output_dir(args.output_dir))
    columns = [c.strip() for c in args.columns.split(",") if c.strip()]
    try:
        dates = [args.date] if args.date else dates_in_range(args.date_from, args.date_to)
        inputs = {d: snapshot_output_path(d, "setupsession_with_outcome", output_dir=out_dir) for d in dates}
    except ValueError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)

    missing = [d for d, p in inputs.items() if not os.path.exists(p)]
    if missing:
        print(f"ERROR: setupsession_with_outcome.csv missing for: {', '.join(missing)} -- "
              "run run_pipeline.py for those dates first (this script never fires Athena)",
              file=sys.stderr)
        sys.exit(1)

    try:
        per_date = {d: filter_rows(read_source(p), columns) for d, p in inputs.items()}
    except ValueError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)

    if args.date:
        out_path = args.out or snapshot_output_path(args.date, ARTIFACT, output_dir=out_dir)
        write_csv(out_path, columns, per_date[args.date])
    else:
        out_path = args.out or os.path.splitext(
            range_output_path(args.date_from, args.date_to, ARTIFACT, output_dir=out_dir))[0] + ".xlsx"
        write_xlsx(out_path, columns, {ddmmyyyy_to_date(d).isoformat(): rows for d, rows in per_date.items()})

    total = sum(len(r) for r in per_date.values())
    print(f"Wrote {total} rows across {len(per_date)} date(s) to {out_path}", file=sys.stderr)
    LOGGER.info("Completed extract_error_report_1xxx_6xxx rows=%d out=%s", total, out_path)


if __name__ == "__main__":
    main()
